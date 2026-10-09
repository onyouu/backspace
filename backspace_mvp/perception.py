"""Local document perception. Transcriptions survive every downstream AI failure.

OCR confidence is a routing heuristic, never a probability that evidence is true.
Native strings, OCR candidates and human edits are separate, immutable observations.
"""
import copy
import hashlib
import importlib.util
import json
import os
import re
from pathlib import Path

VERSION = 'perception-1'
DET = 'PP-OCRv5_mobile_det'
REC = 'korean_PP-OCRv5_mobile_rec'
THRESHOLD = .80
_paddle = None
_surya = None


def file_hash(path):
    h = hashlib.sha256()
    with open(path, 'rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def config():
    from dotenv import load_dotenv
    load_dotenv(override=False)
    threshold = float(os.environ.get('BACKSPACE_OCR_THRESHOLD', THRESHOLD))
    if not 0 <= threshold <= 1:
        raise ValueError('OCR 확인 기준은 0~1이어야 합니다.')
    return {'version': VERSION, 'detector': DET, 'recognizer': REC,
            'threshold': threshold, 'device': 'cpu', 'dpi': 200,
            'surya': os.environ.get('BACKSPACE_SURYA_ENABLED', '0') == '1'}


def readiness():
    return {'paddle_installed': importlib.util.find_spec('paddleocr') is not None,
            'surya_installed': importlib.util.find_spec('surya') is not None,
            'surya_enabled': config()['surya'], 'threshold': config()['threshold'],
            'detector': DET, 'recognizer': REC, 'device': 'cpu'}


class PaddleReader:
    def recognize(self, image):
        global _paddle
        if _paddle is None:
            try:
                from paddleocr import PaddleOCR
            except ImportError:
                raise ValueError('로컬 OCR 설치가 필요합니다. docs/perception.md의 설치 방법을 확인하세요.') from None
            _paddle = PaddleOCR(text_detection_model_name=DET, text_recognition_model_name=REC,
                device='cpu', use_doc_orientation_classify=False, use_doc_unwarping=False,
                use_textline_orientation=False, enable_mkldnn=False, cpu_threads=4)
        import numpy as np
        rows = []
        for result in _paddle.predict(np.asarray(image.convert('RGB')), text_rec_score_thresh=0.0):
            data = result.json
            if isinstance(data, str):
                data = json.loads(data)
            data = data.get('res', data)
            for text, score, poly in zip(data['rec_texts'], data['rec_scores'], data['rec_polys']):
                rows.append({'text': text, 'confidence': float(score),
                             'polygon': [[float(x), float(y)] for x, y in poly]})
        return rows


class SuryaReader:
    """Optional current Surya API; its local backend must be installed separately."""
    def recognize(self, image):
        global _surya
        if _surya is None:
            from surya.inference import SuryaInferenceManager
            from surya.recognition import RecognitionPredictor
            _surya = RecognitionPredictor(SuryaInferenceManager())
        results = _surya([image.convert('RGB')])
        rows = []
        for result in results:
            for block in result.blocks:
                if block.error or block.skipped:
                    continue
                text = re.sub('<[^>]+>', '', block.html)
                import html
                rows.append({'text': html.unescape(text), 'confidence': float(block.confidence),
                             'polygon': block.polygon})
        return rows


def _key(value):
    from .pipeline import digest
    return digest(value)


def _box(poly):
    return [min(p[0] for p in poly), min(p[1] for p in poly),
            max(p[0] for p in poly), max(p[1] for p in poly)]


def _text_key(text):
    return re.sub(r'\s+', '', text).casefold()


def _overlap(a, b):
    area = max(0, min(a[2], b[2]) - max(a[0], b[0])) * max(0, min(a[3], b[3]) - max(a[1], b[1]))
    return area / max(1, min((a[2]-a[0])*(a[3]-a[1]), (b[2]-b[0])*(b[3]-b[1])))


def _order(blocks):
    # Separate clearly disjoint columns; spanning titles precede their columns.
    if not blocks:
        return []
    bounds = sorted({b['bbox'][0] for b in blocks})
    width = max(b['bbox'][2] for b in blocks) - min(bounds)
    cuts = [(b-a, (a+b)/2) for a, b in zip(bounds, bounds[1:]) if b-a > width*.22]
    if cuts:
        cut = max(cuts)[1]
        left = [b for b in blocks if b['bbox'][2] <= cut]
        right = [b for b in blocks if b['bbox'][0] >= cut]
        spans = [b for b in blocks if b not in left and b not in right]
        if left and right and not any(b['bbox'][1] > min(x['bbox'][1] for x in left+right)+20 for b in spans):
            return sorted(spans, key=lambda b:(b['bbox'][1],b['bbox'][0])) + sorted(left,key=lambda b:(b['bbox'][1],b['bbox'][0])) + sorted(right,key=lambda b:(b['bbox'][1],b['bbox'][0]))
    return sorted(blocks, key=lambda b:(round(b['bbox'][1]/5), b['bbox'][0]))


def seal(document):
    """Rebuild the exact, quoteable transcription and offsets after review."""
    accepted = pending = 0
    for page in document['pages']:
        pieces, pos = [], 0
        for order, block in enumerate(_order(page['blocks'])):
            block['order'] = order
            block.pop('start', None)
            block.pop('end', None)
            if block['status'] in ('accepted', 'human_verified') and block['text'].strip():
                block['start'] = pos
                block['end'] = pos + len(block['text'])
                pos = block['end'] + 2
                pieces.append(block['text'])
                accepted += 1
            elif block['status'] == 'needs_review':
                pending += 1
        page['text'] = '\n\n'.join(pieces)
        page['pending_blocks'] = sum(b['status'] == 'needs_review' for b in page['blocks'])
        page['status'] = 'needs_review' if page['pending_blocks'] or page.get('errors') else ('ready' if page['text'] else 'blank' if not page['assets'] else 'visual_only')
    document['coverage'] = {'pages': len(document['pages']), 'accepted_blocks': accepted,
        'pending_blocks': pending, 'pages_with_errors': sum(bool(p.get('errors')) for p in document['pages']),
        'visual_regions': sum(len(p['assets']) for p in document['pages'])}
    document['status'] = 'needs_review' if pending or document['coverage']['pages_with_errors'] else 'ready'
    document['content_hash'] = _key({k:v for k,v in document.items() if k != 'content_hash'})
    return document


def validate(document):
    from jsonschema import Draft202012Validator
    from .model import read_json
    Draft202012Validator(read_json(Path(__file__).parent/'schemas/normalized_material.schema.json')).validate(document)
    ids = set()
    for page in document['pages']:
        for block in page['blocks']:
            if block['id'] in ids:
                raise ValueError('인식 블록 ID가 중복됩니다.')
            ids.add(block['id'])
            if 'start' in block and page['text'][block['start']:block['end']] != block['text']:
                raise ValueError('인식 텍스트 위치가 일치하지 않습니다.')
    if document['content_hash'] != _key({k:v for k,v in document.items() if k != 'content_hash'}):
        raise ValueError('인식 결과 해시가 일치하지 않습니다.')
    return document


def extract_document(path, folder, *, reader=None, rescue=None, progress=None, page_range=None):
    """Persist after EACH page. Failed readers leave native text and diagnosis visible."""
    from .pipeline import save
    from .model import read_json
    from PIL import Image
    path, folder = Path(path), Path(folder)
    if page_range is not None and (not page_range or any(type(n) is not int or n<1 for n in page_range)):
        raise ValueError('페이지 선택은 1 이상의 정수 목록이어야 합니다.')
    sha, settings = file_hash(path), config()
    cache = folder / (sha + '-' + _key([settings, page_range])[:12])
    manifest = cache/'normalized.json'
    if manifest.exists():
        stored = validate(read_json(manifest))
        if stored.get('complete') and not stored['coverage']['pages_with_errors']:
            return stored, manifest
    cache.mkdir(parents=True, exist_ok=True)
    document = {'schema_version':'1.0', 'source_sha256':sha, 'title':path.name,
        'format':path.suffix.lower().lstrip('.'), 'settings':settings, 'revision':1,
        'complete':False, 'pages':[], 'history':[]}
    previous=None
    if manifest.exists():
        # Resume exact completed pages; retry only pages with extraction errors.
        previous = validate(read_json(manifest))
        document['history'], document['revision'] = previous['history'], previous['revision']+1
        done = {p['locator']:p for p in previous['pages'] if not p['errors']}
    else:
        done = {}
    reader = reader or PaddleReader()
    rescue = rescue or (SuryaReader() if settings['surya'] else None)

    def page(locator, width, height):
        return {'locator':locator, 'width':float(width), 'height':float(height),
                'blocks':[], 'assets':[], 'errors':[], 'visual_acknowledged':False}

    def block(p, text, bbox, method, confidence=None, asset=None, candidates=None):
        row = {'id':'pb_'+_key([sha,p['locator'],len(p['blocks']),bbox,method])[:20],
            'raw_text':text, 'text':text, 'bbox':[float(x) for x in bbox], 'method':method,
            'confidence':confidence, 'confidence_basis':'native_not_scored' if method=='native' else 'ocr_self_report',
            'status':'accepted' if method=='native' or confidence is not None and confidence >= settings['threshold'] else 'needs_review',
            'asset_id':asset, 'candidates':candidates or [], 'edits':[]}
        if method=='native' and (not text.strip() or '\ufffd' in text or '\x00' in text):
            row['status'] = 'needs_review'
        p['blocks'].append(row)
        return row

    def image(p, im, bbox, label, visual=True):
        if 'A' in im.getbands():
            from PIL import Image
            background=Image.new('RGBA',im.size,'white')
            im=Image.alpha_composite(background,im.convert('RGBA'))
        im = im.convert('RGB')
        aid = 'pa_'+_key([sha,p['locator'],label,bbox])[:20]
        name = aid+'.png'
        im.save(cache/name)
        p['assets'].append({'id':aid,'file':name,'label':label,'bbox':[float(v) for v in bbox],
            'width':im.width,'height':im.height,'sha256':file_hash(cache/name),
            'visual_review_required':visual})
        try:
            rows = reader.recognize(im)
            for row in rows:
                poly = row['polygon']
                px = _box(poly)
                bw, bh = bbox[2]-bbox[0], bbox[3]-bbox[1]
                mapped = [bbox[0]+px[0]*bw/im.width,bbox[1]+px[1]*bh/im.height,
                          bbox[0]+px[2]*bw/im.width,bbox[1]+px[3]*bh/im.height]
                native_match = next((b for b in p['blocks'] if b['method']=='native' and _text_key(b['text'])==_text_key(row['text']) and _overlap(mapped,b['bbox'])>.5), None)
                candidate = {'engine':'paddle','text':row['text'],'confidence':row['confidence'], 'bbox':mapped, 'asset_id':aid}
                if native_match:
                    native_match['candidates'].append(candidate)
                    continue
                b = block(p,row['text'],mapped,'paddle',row['confidence'],aid,[candidate])
                if b['status']=='needs_review' and rescue:
                    crop = im.crop((max(0,int(px[0])-8),max(0,int(px[1])-8),min(im.width,int(px[2])+8),min(im.height,int(px[3])+8)))
                    try:
                        alternatives = rescue.recognize(crop)
                        b['candidates'].extend({'engine':'surya','text':a['text'],'confidence':a['confidence']} for a in alternatives)
                        # Cross-model scores are incomparable. Human review resolves disagreements.
                        if len(alternatives)==1 and _text_key(alternatives[0]['text'])==_text_key(b['text']) and alternatives[0]['confidence']>=settings['threshold']:
                            b['status']='accepted'
                    except Exception as exc:
                        p['errors'].append('보조 OCR 실행 실패: '+type(exc).__name__)
            if not rows:
                p['assets'][-1]['no_text_detected'] = True
        except Exception as exc:
            p['errors'].append('OCR 실행 실패: '+str(exc))

    def commit(p):
        old=next((v for v in previous['pages'] if v['locator']==p['locator']),None) if previous else None
        if old and p is not old:
            edits={b['id']:b for b in old['blocks'] if b['edits']}
            for b in p['blocks']:
                if b['id'] in edits:
                    b.update({k:copy.deepcopy(edits[b['id']][k]) for k in ('text','status','edits','confidence_basis')})
        document['pages'].append(p)
        save(manifest, validate(seal(document)))
        if progress:
            progress(p['locator'], len(document['pages']))

    suffix = path.suffix.lower()
    if suffix == '.pdf':
        import pdfplumber
        import pypdfium2 as pdfium
        with pdfplumber.open(path) as pdf, pdfium.PdfDocument(path) as renderer:
            for index, native in enumerate(pdf.pages,1):
                if page_range is not None and index not in page_range:
                    continue
                loc=f'p{index}'
                if loc in done:
                    commit(done[loc]); continue
                p=page(loc,native.width,native.height)
                words=native.extract_words()
                lines={}
                for word in words:
                    lines.setdefault(round(word['top']/3),[]).append(word)
                for words_in_line in lines.values():
                    words_in_line.sort(key=lambda w:w['x0'])
                    # Split large gaps so two columns never become one invented sentence.
                    groups=[[]]
                    for w in words_in_line:
                        if groups[-1] and w['x0']-groups[-1][-1]['x1']>max(25,native.width*.06):
                            groups.append([])
                        groups[-1].append(w)
                    for group in groups:
                        block(p,' '.join(w['text'] for w in group),[group[0]['x0'],min(w['top'] for w in group),group[-1]['x1'],max(w['bottom'] for w in group)],'native')
                rendered=renderer[index-1]
                regions=[]
                for picture in native.images:
                    bbox=[max(0,picture['x0']),max(0,picture['top']),min(native.width,picture['x1']),min(native.height,picture['bottom'])]
                    if bbox[2]>bbox[0] and bbox[3]>bbox[1] and not any(_overlap(bbox,r)>.98 for r in regions):
                        regions.append(bbox)
                if not regions:
                    # Vectorized glyphs can look like text while exposing no native characters.
                    if native.curves or native.rects or native.lines:
                        regions=[[0,0,native.width,native.height]]
                for n,bbox in enumerate(regions,1):
                    crop=(bbox[0],native.height-bbox[3],native.width-bbox[2],bbox[1])
                    bitmap=rendered.render(scale=settings['dpi']/72,crop=crop)
                    image(p,bitmap.to_pil(),bbox,f'page_{index}_image_{n}')
                    bitmap.close()
                rendered.close()
                commit(p)
    elif suffix=='.pptx':
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from io import BytesIO
        deck=Presentation(path)
        def walk(p, shapes, transform=(1,1,0,0)):
            sx,sy,tx,ty=transform
            for shape in shapes:
                bbox=[(shape.left*sx+tx)/12700,(shape.top*sy+ty)/12700,
                      ((shape.left+shape.width)*sx+tx)/12700,((shape.top+shape.height)*sy+ty)/12700]
                if shape.shape_type==MSO_SHAPE_TYPE.GROUP:
                    xfrm=shape._element.grpSpPr.xfrm
                    gx=float(xfrm.ext.cx)/max(1,float(xfrm.chExt.cx)); gy=float(xfrm.ext.cy)/max(1,float(xfrm.chExt.cy))
                    walk(p,shape.shapes,(sx*gx,sy*gy,tx+sx*(xfrm.off.x-gx*xfrm.chOff.x),ty+sy*(xfrm.off.y-gy*xfrm.chOff.y)))
                elif shape.has_table:
                    for ri,row in enumerate(shape.table.rows):
                        for ci,cell in enumerate(row.cells):
                            if cell.text.strip():
                                x=bbox[0]+sum(c.width for c in list(shape.table.columns)[:ci])*sx/12700
                                y=bbox[1]+sum(r.height for r in list(shape.table.rows)[:ri])*sy/12700
                                block(p,cell.text,[x,y,x+shape.table.columns[ci].width*sx/12700,y+row.height*sy/12700],'native')
                elif shape.has_text_frame and shape.text.strip():
                    block(p,shape.text,bbox,'native')
                elif shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
                    im=Image.open(BytesIO(shape.image.blob)).convert('RGBA')
                    im=im.crop((round(im.width*shape.crop_left),round(im.height*shape.crop_top),round(im.width*(1-shape.crop_right)),round(im.height*(1-shape.crop_bottom))))
                    if shape.rotation:
                        im=im.rotate(-shape.rotation,expand=True)
                    image(p,im,bbox,f'{p["locator"]}_image_{shape.shape_id}')
                elif getattr(shape,'has_chart',False):
                    p['errors'].append('차트 의미는 원본 확인이 필요합니다. PDF로 내보내면 시각 영역을 보존할 수 있습니다.')
        for index,slide in enumerate(deck.slides,1):
            if page_range is not None and index not in page_range:
                continue
            loc=f'slide{index}'
            if loc in done:
                commit(done[loc]); continue
            p=page(loc,deck.slide_width/12700,deck.slide_height/12700)
            walk(p,slide.shapes)
            # Backgrounds can be inherited from a layout/master.
            for owner in (slide,slide.slide_layout,slide.slide_layout.slide_master):
                backgrounds=owner._element.xpath('.//p:bg//a:blip')
                if backgrounds:
                    for n,blip in enumerate(backgrounds):
                        rid=blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                        if rid:
                            image(p,Image.open(BytesIO(owner.part.related_part(rid).blob)),[0,0,p['width'],p['height']],f'{loc}_background_{n}')
                    break
            commit(p)
    elif suffix in ('.png','.jpg','.jpeg'):
        from PIL import ImageOps
        im=ImageOps.exif_transpose(Image.open(path))
        p=page('p1',im.width,im.height)
        image(p,im,[0,0,im.width,im.height],'image_1')
        commit(p)
    elif suffix in ('.txt','.md'):
        p=page('p1',1,1)
        text=path.read_text(encoding='utf-8-sig')
        if text.strip():
            block(p,text,[0,0,1,1],'native')
        commit(p)
    else:
        raise ValueError('PDF/PPTX/PNG/JPG/UTF-8 TXT/MD 입력을 지원합니다.')
    document['complete']=True
    save(manifest,validate(seal(document)))
    return document,manifest


def review_document(manifest, revision, block_id=None, text=None, locator=None, acknowledge_visual=False):
    from .model import read_json
    from .pipeline import save
    doc=validate(read_json(manifest))
    if type(revision) is not int or doc['revision']!=revision:
        raise ValueError('인식 결과가 변경됐습니다. 새로 열어 다시 확인하세요.')
    event={'revision':revision,'block_id':block_id,'locator':locator}
    if block_id:
        row=next((b for p in doc['pages'] for b in p['blocks'] if b['id']==block_id),None)
        if row is None or not isinstance(text,str) or len(text)>8000:
            raise ValueError('수정할 블록과 8000자 이하 텍스트가 필요합니다.')
        event.update(previous=row['text'],text=text)
        row['edits'].append(copy.deepcopy(event))
        row.update(text=text,status='human_verified' if text.strip() else 'excluded')
        row['confidence_basis']='human_verified_transcription'
    elif acknowledge_visual and locator:
        page=next((p for p in doc['pages'] if p['locator']==locator),None)
        if not page:
            raise ValueError('페이지를 찾을 수 없습니다.')
        page['visual_acknowledged']=True
        event['visual_acknowledged']=True
    else:
        raise ValueError('확인할 블록 또는 시각 영역이 필요합니다.')
    doc['history'].append(event)
    doc['revision']+=1
    save(manifest,validate(seal(doc)))
    return doc


def source_from_document(path, document, kind, exam_id=None, material_type=None):
    """Existing graph contracts consume exact, accepted transcription strings."""
    sha=document['source_sha256']
    blocked=[]
    excerpts=[]
    for page in document['pages']:
        unsafe=bool(page['errors']) or (kind=='exam' and (page['pending_blocks'] or page['assets'] and not page['visual_acknowledged']))
        if unsafe:
            blocked.append(page['locator'])
        if page['text'].strip() and not unsafe:
            excerpts.append({'locator':page['locator'],'text':page['text']})
    source={'id':'src_'+_key([sha,kind,exam_id])[:20], 'kind':kind,'title':Path(path).name,
        'path':str(Path(path).resolve()),'sha256':sha,'format':'txt' if document['format']=='md' else document['format'],
        'extraction_status':'ai_transcribed_needs_review' if document['coverage']['visual_regions'] else 'text_extracted_needs_review',
        'excerpts':excerpts,'page_count':max(1,len(document['pages'])),
        'perception':{'version':VERSION,'content_hash':document['content_hash'],'revision':document['revision'],
            'coverage':document['coverage'],'blocked_pages':blocked,
            'blocks':[{'locator':p['locator'],**{k:b[k] for k in ('id','start','end','bbox','method','confidence','confidence_basis','status','asset_id','candidates','selected_source','selected_text','initial_text') if k in b}}
                for p in document['pages'] for b in p['blocks'] if 'start' in b and p['locator'] not in blocked]}}
    if kind in ('exam','answer_key'):
        if not exam_id:
            raise ValueError('--exam-id가 필요합니다.')
        source['exam_id']=exam_id
    elif material_type:
        source['material_type']=material_type
    return source


def transcription_refs(source, locator, start, end):
    return [copy.deepcopy(b) for b in source.get('perception',{}).get('blocks',[])
        if b['locator']==locator and b['start']<end and start<b['end']]
