import copy
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from backspace_mvp.materials import export_material_graph, validate_material_graph
from backspace_mvp.model import read_json
from backspace_mvp.pipeline import ingest, save
from backspace_mvp.nvidia import NvidiaProvider


class MaterialFixture:
    """A deterministic extraction fixture; does not pretend to be Jev or an API result."""
    calls = 0

    def complete(self, prompt, schema):
        self.calls += 1
        if 'verdicts' in schema['properties']:
            return {'verdicts': [{'index': 0, 'status': 'supported', 'reason': '고정 검증 fixture'}]}
        text = json.JSONDecoder().raw_decode(prompt.split('\n', 1)[1])[0]['text']
        concepts = []
        for word in ('산소', 'ATP', '세포호흡'):
            if word in text:
                concepts.append({'label': word, 'surface_form': word, 'quote': text, 'confidence': 0.75,
                                 'explicit_emphasis': [text] if word == '산소' and '중요' in text else []})
        relations = [{'from_label': '산소', 'to_label': 'ATP', 'relation': 'related_to',
                      'statement': '산소 조건과 ATP 생성을 비교한다.', 'conditions': ['산소 공급 조건의 비교'],
                      'quote': text, 'confidence': 0.6}] if 'ATP' in text else []
        return {'concepts': concepts, 'relations': relations}


class MaterialGraphTests(unittest.TestCase):
    def test_cumulative_pages_types_standalone_concept_and_replay(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            deck = Presentation()
            for content in ('중요: 산소와 ATP 생성의 관계를 비교한다.', '산소와 ATP 생성의 조건을 다시 비교한다.'):
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = content
            ppt = root / 'material.pptx'
            deck.save(ppt)
            worksheet = root / 'worksheet.txt'
            worksheet.write_text('산소와 ATP 생성을 비교한다. 세포호흡.', encoding='utf-8')
            state = root / 'dataset.json'
            provider = MaterialFixture()
            first = ingest(state, ppt, 'learning_material', provider=provider)
            graph_path = root / 'material_graph.json'
            self.assertEqual(read_json(graph_path)['revision'], 1)
            data = ingest(state, worksheet, 'learning_material', provider=provider, material_type='worksheet')
            graph = read_json(graph_path)
            oxygen = next(c for c in graph['concepts'] if c['label'] == '산소')
            self.assertEqual(oxygen['frequency'], 3)
            self.assertEqual(oxygen['source_count'], 2)
            self.assertEqual(oxygen['source_types'], ['ppt', 'worksheet'])
            self.assertEqual(oxygen['explicit_emphasis']['frequency'], 1)
            self.assertEqual(len(oxygen['extraction_confidences']), 3)
            self.assertEqual(next(c for c in graph['concepts'] if c['label'] == '세포호흡')['frequency'], 1)
            self.assertEqual(graph['relations'][0]['frequency'], 3)
            self.assertTrue(all(a['conditions'] for a in graph['relations'][0]['assertions']))
            self.assertEqual(graph['revision'], 2)
            self.assertEqual(len(list((root / 'material_graph_history').glob('*.json'))), 1)
            before = graph_path.read_bytes()
            calls = provider.calls
            ingest(state, worksheet, 'learning_material', provider=provider, material_type='worksheet')
            self.assertEqual(provider.calls, calls)
            self.assertEqual(graph_path.read_bytes(), before)
            validate_material_graph(graph, data)
            for concept in graph['concepts']:
                for ref in concept['evidence']:
                    source = next(s for s in data['sources'] if s['id'] == ref['source'])
                    text = next(p['text'] for p in source['excerpts'] if p['locator'] == ref['locator'])
                    self.assertEqual(text[ref['char_start']:ref['char_end']], ref['raw_evidence'])
                    self.assertIsNone(ref['question_id'])
            # Material aggregation does not depend on unused exam-only concepts in the registry.
            exam_catalog = copy.deepcopy(data)
            exam_catalog['concepts'].append({'id': 'exam_only', 'label': '시험 전용', 'aliases': []})
            export_material_graph(exam_catalog, graph_path)
            self.assertEqual(graph_path.read_bytes(), before)

    def test_refresh_replaces_evidence_and_keeps_revision_on_same_result(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            material = root / 'material.txt'
            material.write_text('산소와 ATP 생성을 비교한다.', encoding='utf-8')
            state = root / 'dataset.json'
            ingest(state, material, 'learning_material', provider=MaterialFixture(), material_type='worksheet')
            graph_path = root / 'material_graph.json'
            before = graph_path.read_bytes()
            data = ingest(state, material, 'learning_material', provider=MaterialFixture(), refresh=True)
            self.assertEqual(data['sources'][0]['material_type'], 'worksheet')
            self.assertEqual(len(data['material_observations']), 2)
            self.assertEqual(len(data['material_relations']), 1)
            self.assertEqual(graph_path.read_bytes(), before)

    def test_failed_quote_or_invented_emphasis_preserves_state(self):
        class Corrupted(MaterialFixture):
            def complete(self, prompt, schema):
                value = super().complete(prompt, schema)
                if 'concepts' in value:
                    value['concepts'][0]['explicit_emphasis'] = ['산소와 ATP 생성을 비교한다.']
                return value
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            material = root / 'material.txt'
            material.write_text('산소와 ATP 생성을 비교한다.', encoding='utf-8')
            state = root / 'dataset.json'
            ingest(state, material, 'learning_material', provider=MaterialFixture())
            before = state.read_bytes()
            with self.assertRaises(ValueError):
                ingest(state, material, 'learning_material', provider=Corrupted(), refresh=True)
            self.assertEqual(state.read_bytes(), before)

    def test_graph_and_observation_tampering_rejected(self):
        from backspace_mvp.model import validate
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            material = root / 'material.txt'
            material.write_text('산소.', encoding='utf-8')
            data = ingest(root / 'dataset.json', material, 'learning_material', provider=MaterialFixture())
            graph = read_json(root / 'material_graph.json')
            graph['concepts'][0]['frequency'] += 1
            with self.assertRaisesRegex(ValueError, 'dataset'):
                validate_material_graph(graph, data)
            corrupted = copy.deepcopy(data)
            corrupted['material_observations'][0]['evidence'][0]['quote'] = '없는 원문'
            with self.assertRaises(ValueError):
                validate(corrupted)
            corrupted = copy.deepcopy(data)
            corrupted['material_observations'].append(copy.deepcopy(corrupted['material_observations'][0]))
            with self.assertRaises(ValueError):
                validate(corrupted)

    def test_usage_is_deduplicated_and_credentials_absent(self):
        class Response:
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def read(self):
                return json.dumps({'usage': {'prompt_tokens': 3, 'completion_tokens': 2},
                    'choices': [{'finish_reason': 'stop', 'message': {'content': '{"ok":true}'}}]}).encode()
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            with patch.dict(os.environ, {'NVIDIA_API_KEY': 'secret-test-only', 'NVIDIA_LOG_DIR': str(root / 'calls')}, clear=True):
                provider = NvidiaProvider(env_file=root / 'absent.env', usage_path=root / 'usage_log.json')
                with patch('urllib.request.urlopen', return_value=Response()):
                    provider.complete('input', {'type': 'object'})
                text = (root / 'usage_log.json').read_text(encoding='utf-8')
                self.assertNotIn('secret-test-only', text)
                log = json.loads(text)
                self.assertEqual(len(log['runs']), 1)
                self.assertEqual(len(next(iter(log['runs'].values()))['attempts']), 1)
                self.assertIsNone(next(iter(log['runs'].values()))['cost'])

    def test_legacy_graph_marks_missing_observations_without_inventing_confidence(self):
        sample = Path(__file__).resolve().parents[1] / 'data/sample.json'
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            data = ingest(root / 'dataset.json', sample, 'learning_material')
            graph = read_json(root / 'material_graph.json')
            self.assertTrue(graph['needs_reextraction'])
            self.assertTrue(all(c['observation_status'] == 'legacy_needs_reextraction' for c in graph['concepts']))
            self.assertTrue(all(e['confidence'] is None for c in graph['concepts'] for e in c['evidence']))
            validate_material_graph(graph, data)
