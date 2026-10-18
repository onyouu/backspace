import copy
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import networkx as nx
from backspace_mvp.analysis import analyze, mine
from backspace_mvp.model import read_json, validate
from backspace_mvp.pipeline import ingest, build, generate_plan, save, read_source
from backspace_mvp.nvidia import NvidiaProvider

SAMPLE = Path(__file__).resolve().parents[1] / 'data/sample.json'


class PipelineTests(unittest.TestCase):
    def test_pptx_and_pdf_readers(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = '세포호흡'
            deck.save(path / 'lesson.pptx')
            source = read_source(path / 'lesson.pptx', 'learning_material')
            self.assertEqual(source['excerpts'][0]['text'], '세포호흡')
            self.assertEqual(source['format'], 'pptx')
            pdf = PdfWriter()
            pdf.add_blank_page(width=100, height=100)
            pdf.write(path / 'scan.pdf')
            blank = read_source(path / 'scan.pdf', 'exam', 'test')
            self.assertEqual(blank['excerpts'], [])
            self.assertEqual(blank['perception']['coverage']['pages'], 1)

    def test_raw_source_provider_and_idempotence(self):
        from backspace_mvp.pipeline import empty_dataset
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)
            text = path / 'lesson.txt'
            text.write_text('세포호흡은 에너지를 전환한다.', encoding='utf-8')
            source = read_source(text, 'learning_material')
            class Provider:
                calls = 0
                def complete(self, prompt, schema):
                    self.calls += 1
                    result = empty_dataset()
                    result['sources'] = [source]
                    result['concepts'] = [{'id': 'respiration', 'label': '세포호흡', 'aliases': []}]
                    return {'concepts': [{'label': '세포호흡', 'surface_form': '세포호흡',
                        'quote': '세포호흡은 에너지를 전환한다.', 'confidence': 0.8, 'explicit_emphasis': []}], 'relations': []}
            provider = Provider()
            state = path / 'state.json'
            ingest(state, text, 'learning_material', provider=provider)
            ingest(state, text, 'learning_material', provider=provider)
            self.assertEqual(provider.calls, 1)
            self.assertEqual(read_json(state)['sources'][0]['sha256'], source['sha256'])

    def test_complete_flow_and_resume(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            state, model, plan = [root / f'{n}.json' for n in ('dataset', 'model', 'plan')]
            ingest(state, SAMPLE, 'learning_material')
            self.assertEqual(read_json(state)['questions'], [])
            ingest(state, SAMPLE, 'exam')
            before = state.read_bytes()
            ingest(state, SAMPLE, 'exam')
            self.assertEqual(before, state.read_bytes())
            report = build(state, model)
            self.assertEqual(report['counts']['questions'], 6)
            first = generate_plan(model, plan)
            first['tasks'][0]['status'] = 'completed'
            save(plan, first)
            second = generate_plan(model, plan, feedback='추가 데이터 대기')
            self.assertEqual(second['revision'], 2)
            self.assertEqual(second['tasks'][0]['status'], 'completed')
            self.assertEqual(len(second['events']), 2)

    def test_no_cross_question_path(self):
        a, b = nx.MultiDiGraph(), nx.MultiDiGraph()
        for g, codes in ((a, ['compare', 'infer_cause']), (b, ['infer_cause', 'justify'])):
            for i, code in enumerate(codes):
                g.add_node(i, kind='reasoning', code=code)
            g.add_edge(0, 1, relation='precedes')
        result = mine({'a': a, 'b': b}, [{'id': 'a', 'exam_id': 'e', 'evidence': []}, {'id': 'b', 'exam_id': 'e', 'evidence': []}])
        self.assertEqual(result['paths'], [])

    def test_signature_support(self):
        report = analyze(read_json(SAMPLE))[0]
        diamond = [r for r in report['signature']['subgraphs'] if len(r['pattern'][0]) == 4]
        self.assertEqual(diamond[0]['question_ids'], ['Q5', 'Q6'])
        for rows in report['signature'].values():
            for row in rows:
                self.assertEqual(row['support'], len(set(row['question_ids'])))
                if row['support'] == 1:
                    self.assertFalse(row['is_repeated'])

    def test_evidence_rejected(self):
        data = read_json(SAMPLE)
        data['questions'][0]['intent']['evidence'][0]['quote'] = '없는 인용'
        with self.assertRaises(ValueError):
            validate(data)

    def test_reference_write_rejected(self):
        with self.assertRaises(ValueError):
            save(Path('sources/result.json'), {})

    def test_model_corruption_does_not_commit(self):
        class BadProvider:
            def complete(self, prompt, schema):
                result = read_json(SAMPLE)
                return result
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)
            text = path / 'lesson.txt'
            text.write_text('세포호흡', encoding='utf-8')
            state = path / 'state.json'
            from jsonschema import ValidationError
            with self.assertRaises((ValueError, ValidationError)):
                ingest(state, text, 'learning_material', provider=BadProvider())
            self.assertFalse(state.exists())

    def test_nvidia_contract(self):
        class Response:
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def read(self): return json.dumps({'choices': [{'finish_reason': 'stop', 'message': {'content': '{"ok": true}'}}]}).encode()
        with patch.dict(os.environ, {'NVIDIA_API_KEY': 'fake-test-key'}, clear=True):
            with patch('urllib.request.urlopen', return_value=Response()) as call:
                result = NvidiaProvider(env_file=Path('not-present.env')).complete('test', {'type': 'object'})
                payload = json.loads(call.call_args.args[0].data)
                self.assertEqual(payload['model'], NvidiaProvider.SUPER)
                self.assertNotIn('fake-test-key', json.dumps(payload))
                self.assertTrue(result['ok'])

    def test_missing_key(self):
        with patch.dict(os.environ, {}, clear=True):
            with self.assertRaises(ValueError):
                NvidiaProvider(env_file=Path('not-present.env'))

    def test_review_reasoning_override_does_not_change_extraction(self):
        class Response:
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def read(self): return b'{"choices":[{"finish_reason":"stop","message":{"content":"{}"}}]}'
        with patch.dict(os.environ, {'NVIDIA_API_KEY': 'fake-test', 'NVIDIA_REASONING_EFFORT': 'low',
                                     'NVIDIA_REVIEW_REASONING_EFFORT': 'none'}, clear=True):
            exam = NvidiaProvider(env_file=Path('not-present.env'), task='exam')
            review = NvidiaProvider(env_file=Path('not-present.env'), task='review')
            self.assertEqual(exam.reasoning_effort, 'low')
            self.assertEqual(review.reasoning_effort, 'none')
            with patch('urllib.request.urlopen', return_value=Response()) as call:
                review.complete('review', {'type': 'object'})
                request = json.loads(call.call_args.args[0].data)
                self.assertEqual(request['reasoning_effort'], 'none')
                self.assertNotIn('reasoning_budget', request)

    def test_timeout_diagnostics_do_not_contain_credentials(self):
        with tempfile.TemporaryDirectory() as folder:
            with patch.dict(os.environ, {'NVIDIA_API_KEY':'secret-test-key', 'NVIDIA_LOG_DIR':folder}, clear=True):
                provider = NvidiaProvider(env_file=Path('missing.env'))
                with patch('urllib.request.urlopen', side_effect=TimeoutError()):
                    with self.assertRaisesRegex(ValueError, 'timeout'):
                        provider.complete('test', {'type':'object'})
                log = provider.log_path.read_text(encoding='utf-8')
                self.assertNotIn('secret-test-key', log)
                self.assertEqual([a['status'] for a in json.loads(log)['attempts']], ['timeout'])

    def test_task_routing_and_credentials(self):
        with patch.dict(os.environ, {'NVIDIA_SUPER_API_KEY': 'super-test', 'NVIDIA_LIGHTNING_API_KEY': 'lightning-test'}, clear=True):
            material = NvidiaProvider(env_file=Path('missing.env'), task='learning_material')
            exam = NvidiaProvider(env_file=Path('missing.env'), task='exam')
            self.assertEqual(material.model, NvidiaProvider.LIGHTNING)
            self.assertEqual(material.key, 'lightning-test')
            self.assertEqual(exam.model, NvidiaProvider.SUPER)
            self.assertEqual(exam.key, 'super-test')

    def test_timeout_does_not_resend_same_prompt_to_fallback(self):
        class Response:
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def read(self): return b'{"choices":[{"finish_reason":"stop","message":{"content":"{}"}}]}'
        with patch.dict(os.environ, {'NVIDIA_API_KEY': 'fake-test', 'NVIDIA_MATERIAL_MODEL': NvidiaProvider.LIGHTNING}, clear=True):
            provider = NvidiaProvider(env_file=Path('missing.env'), task='learning_material')
            with patch('urllib.request.urlopen', side_effect=[TimeoutError(), Response()]) as call:
                with self.assertRaisesRegex(ValueError,'timeout'):
                    provider.complete('test', {'type':'object'})
                self.assertEqual(call.call_count,1)
                self.assertIsNone(provider.last_model)
                first = json.loads(call.call_args_list[0].args[0].data)
                self.assertFalse(first['chat_template_kwargs']['enable_thinking'])

    def test_auth_failure_is_not_retried(self):
        from urllib.error import HTTPError
        with patch.dict(os.environ, {'NVIDIA_API_KEY': 'fake-test'}, clear=True):
            with patch('urllib.request.urlopen', side_effect=HTTPError('https://example.com',401,'Unauthorized',{},None)) as call:
                with self.assertRaises(ValueError):
                    NvidiaProvider(env_file=Path('missing.env')).complete('test', {'type':'object'})
                self.assertEqual(call.call_count, 1)

    def test_review_invalid_output_waits_for_explicit_stage_resume(self):
        class Response:
            def __init__(self, content): self.content = content
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def read(self): return json.dumps({'choices': [{'finish_reason': 'stop', 'message': {'content': self.content}}]}).encode()
        with patch.dict(os.environ, {'NVIDIA_API_KEY': 'fake-test', 'NVIDIA_MATERIAL_MODEL': NvidiaProvider.LIGHTNING}, clear=True):
            provider = NvidiaProvider(env_file=Path('missing.env'), task='review')
            self.assertEqual(provider.fallback, provider.model)
            with patch('urllib.request.urlopen', side_effect=[Response('bad-json'), Response('{"ok":true}')]) as call:
                with self.assertRaises(ValueError):provider.complete('review', {'type': 'object', 'required': ['ok']})
                models = [json.loads(c.args[0].data)['model'] for c in call.call_args_list]
                self.assertEqual(models, [NvidiaProvider.SUPER])

    def test_schema_failure_does_not_swap_model_with_identical_payload(self):
        class Response:
            def __init__(self, content): self.content = content
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def read(self): return json.dumps({'choices':[{'finish_reason':'stop','message':{'content':self.content}}]}).encode()
        with patch.dict(os.environ, {'NVIDIA_API_KEY':'fake-test', 'NVIDIA_MATERIAL_MODEL': NvidiaProvider.LIGHTNING}, clear=True):
            provider = NvidiaProvider(env_file=Path('missing.env'), task='learning_material')
            with patch('urllib.request.urlopen', side_effect=[Response('{}'), Response('{"ok":true}')]) as call:
                with self.assertRaises(ValueError):provider.complete('test', {'type':'object','required':['ok']})
                self.assertEqual(call.call_count, 1)
                self.assertIsNone(provider.last_model)
