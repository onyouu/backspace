"""Perception/state contracts with fake OCR, never remote models or providers."""
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from PIL import Image
from backspace_mvp.perception import extract_document, review_document, source_from_document, validate
from backspace_mvp.pipeline import ingest
from backspace_mvp.grounded import reconcile_material
from backspace_mvp.web import WebStore, MAX_FILE


class Reader:
    def __init__(self, score=.95, text='ATP', fail=False):
        self.score,self.text,self.fail,self.calls=score,text,fail,0
    def recognize(self,image):
        self.calls+=1
        if self.fail:
            raise ValueError('local reader unavailable')
        return [{'text':self.text,'confidence':self.score,'polygon':[[1,1],[30,1],[30,12],[1,12]]}]


class PerceptionTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name); self.folder=self.root/'perception'
        self.block=patch('urllib.request.urlopen',side_effect=AssertionError('No remote API in tests'))
        self.block.start(); self.addCleanup(self.block.stop)
    def picture(self):
        p=self.root/'photo.png'; Image.new('RGB',(100,40),'white').save(p); return p

    def test_low_confidence_is_saved_and_human_edit_is_versioned(self):
        doc,p=extract_document(self.picture(),self.folder,reader=Reader(.41,'AT?'))
        b=doc['pages'][0]['blocks'][0]
        self.assertEqual(doc['pages'][0]['text'],'')
        self.assertEqual(doc['coverage']['pending_blocks'],1)
        updated=review_document(p,1,b['id'],'ATP')
        self.assertEqual(updated['pages'][0]['text'],'ATP')
        self.assertEqual(updated['pages'][0]['blocks'][0]['raw_text'],'AT?')
        self.assertEqual(updated['revision'],2)
        self.assertEqual(updated['pages'][0]['blocks'][0]['confidence'],.41)
        with self.assertRaises(ValueError):review_document(p,1,b['id'],'ATP')
        validate(updated)

    def test_ocr_cache_is_reused_without_reader_calls(self):
        reader=Reader(); image=self.picture()
        a,_=extract_document(image,self.folder,reader=reader)
        b,_=extract_document(image,self.folder,reader=reader)
        self.assertEqual(reader.calls,1);self.assertEqual(a,b)

    def test_exam_visual_page_is_withheld_until_review(self):
        image=self.picture();doc,p=extract_document(image,self.folder,reader=Reader())
        source=source_from_document(image,doc,'exam','exam1')
        self.assertEqual(source['excerpts'],[])
        self.assertEqual(source['perception']['blocked_pages'],['p1'])
        updated=review_document(p,1,locator='p1',acknowledge_visual=True)
        source=source_from_document(image,updated,'exam','exam1')
        self.assertEqual(source['excerpts'][0]['text'],'ATP')

    def test_failed_ocr_leaves_original_and_native_text(self):
        image=self.picture();doc,p=extract_document(image,self.folder,reader=Reader(fail=True))
        self.assertTrue(p.exists());self.assertEqual(doc['coverage']['pages_with_errors'],1)
        self.assertTrue((p.parent/doc['pages'][0]['assets'][0]['file']).exists())
        fixed,_=extract_document(image,self.folder,reader=Reader())
        self.assertEqual(fixed['coverage']['pages_with_errors'],0)
        self.assertEqual(fixed['revision'],2)

    def test_surya_disagreement_does_not_auto_accept_high_score(self):
        doc,_=extract_document(self.picture(),self.folder,reader=Reader(.4,'AT?'),rescue=Reader(.99,'ATP'))
        b=doc['pages'][0]['blocks'][0]
        self.assertEqual(b['status'],'needs_review');self.assertEqual(len(b['candidates']),2)

    def test_native_text_survives_llm_failure(self):
        p=self.root/'lesson.txt';p.write_text('ATP 생성에는 양성자 기울기가 필요하다.',encoding='utf-8')
        class Failure:
            def complete(self,*a):raise ValueError('LLM failure')
        with self.assertRaisesRegex(ValueError,'LLM failure'):
            ingest(self.root/'dataset.json',p,'learning_material',provider=Failure())
        manifests=list(self.folder.glob('*/normalized.json'))
        self.assertEqual(len(manifests),1);self.assertFalse((self.root/'dataset.json').exists())
        self.assertIn('ATP',manifests[0].read_text(encoding='utf-8'))

    def test_ppt_table_group_and_embedded_picture_are_read(self):
        from pptx import Presentation
        from pptx.util import Inches
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
        group=slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text='native heading'
        table=slide.shapes.add_table(1,2,Inches(1),Inches(3),Inches(3),Inches(1)).table
        table.cell(0,0).text='cell A';table.cell(0,1).text='cell B'
        slide.shapes.add_picture(str(self.picture()), Inches(5), Inches(1))
        path=self.root/'deck.pptx';deck.save(path)
        r=Reader();doc,_=extract_document(path,self.folder,reader=r)
        text=doc['pages'][0]['text']
        for word in ('native heading','cell A','cell B','ATP'):self.assertIn(word,text)
        self.assertEqual(r.calls,1)

    def test_pdf_native_plus_image_and_blank_page(self):
        from reportlab.pdfgen.canvas import Canvas
        path=self.root/'mixed.pdf';canvas=Canvas(str(path))
        canvas.drawString(70,740,'native heading');canvas.drawImage(str(self.picture()),70,600,width=200,height=80)
        canvas.showPage();canvas.showPage();canvas.save()
        doc,_=extract_document(path,self.folder,reader=Reader())
        self.assertIn('native heading',doc['pages'][0]['text'])
        self.assertIn('ATP',doc['pages'][0]['text'])
        self.assertEqual(doc['pages'][1]['status'],'blank')
        self.assertEqual(doc['coverage']['pages'],2)

    def test_exact_quote_reconciliation_updates_edges_and_does_not_invent_quote(self):
        quote='사람들은 작은 상처는 오래 간직하고, 큰 은혜는 얼른 망각해버린다'
        old='사람들이 작은 상처를 오래 간직하는 것'
        value={'concepts':[{'label':old,'surface_form':'사람들이 작은 상처를 오래 간직한다','quote':quote,'confidence':.9,'explicit_emphasis':[]},
            {'label':'ATP','surface_form':'ATP','quote':'not in source','confidence':.9,'explicit_emphasis':[]}],
            'relations':[{'from_label':old,'to_label':'ATP'}]}
        changes=reconcile_material(value,{'text':quote},[])
        self.assertEqual(value['concepts'][0]['surface_form'],quote)
        self.assertEqual(value['relations'][0]['from_label'],quote)
        self.assertEqual(value['concepts'][1]['quote'],'not in source')
        self.assertEqual(len(changes),1)

    def test_100mb_stream_boundary_and_interrupted_cleanup(self):
        store=WebStore(storage=self.root/'web')
        class Chunks:
            def __init__(self):self.largest=0
            def read(self,n):self.largest=max(self.largest,n);return b'x'*n
        chunks=Chunks()
        row=store.upload_stream(chunks,MAX_FILE,'large.txt','teacher','material')
        self.assertEqual(row['bytes'],100*1024*1024);self.assertLessEqual(chunks.largest,1024*1024)
        with self.assertRaises(ValueError):store.upload_stream(io.BytesIO(b'x'),MAX_FILE+1,'large.txt','teacher','material')
        with self.assertRaisesRegex(ValueError,'전송'):store.upload_stream(io.BytesIO(b'x'),2,'short.txt','teacher','material')
        self.assertFalse(list((store.workspace/'uploads').glob('*.upload')))

    def test_page_range_cache_does_not_replace_full_document(self):
        from reportlab.pdfgen.canvas import Canvas
        path=self.root/'pages.pdf';canvas=Canvas(str(path))
        canvas.drawString(70,700,'page one');canvas.showPage();canvas.drawString(70,700,'page two');canvas.save()
        subset,sp=extract_document(path,self.folder,reader=Reader(),page_range=[2])
        full,fp=extract_document(path,self.folder,reader=Reader())
        self.assertNotEqual(sp,fp);self.assertEqual(len(subset['pages']),1);self.assertEqual(len(full['pages']),2)
